"""Génère supabase/002_artaluz_seed.sql à partir de la grille tarifaire Excel.

Usage : python3 scripts/build_seed.py <grille.xlsx> supabase/002_artaluz_seed.sql
Relancer à chaque nouvelle version de la grille : le seed est idempotent (upsert).
"""
import sys
from openpyxl import load_workbook

SUPPORT = [("Bâche", "bache"), ("Sticker", "sticker"), ("Magnet", "magnet"),
           ("Poster", "poster"), ("Tableau alu", "alu")]


def q(s):
    return "null" if s is None else "'" + str(s).replace("'", "''") + "'"


def cents(x):
    return int(round(float(x) * 100))


def support_of(label):
    for prefix, code in SUPPORT:
        if str(label).startswith(prefix):
            return code
    raise ValueError(f"Support inconnu : {label}")


def main(xlsx, out):
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Produits"]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    need = ["Réf", "Support", "Format", "Largeur (cm)", "Hauteur (cm)", "PRHT",
            "Prix TTC final (€)", "Qté minimum"]
    missing = [h for h in need if h not in col]
    if missing:
        sys.exit(f"Colonnes manquantes dans la grille : {missing}")

    rows, seen = [], set()
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        ref, label = r[col["Réf"]], r[col["Support"]]
        if not ref or not label or r[col["Prix TTC final (€)"]] is None:
            continue
        if ref in seen:
            sys.exit(f"Référence en double : {ref}")
        seen.add(ref)
        rows.append("(%s,%s,%s,%s,%s,%s,%d,%d,%d,%d)" % (
            q(ref), q(support_of(label)), q(label), q(r[col["Format"]]),
            r[col["Largeur (cm)"]], r[col["Hauteur (cm)"]],
            cents(r[col["Prix TTC final (€)"]]), cents(r[col["PRHT"]] or 0),
            int(r[col["Qté minimum"]] or 1), i))

    fin = []
    wf = wb["Finitions"]
    fh = [c.value for c in wf[1]]
    fc = {h: i for i, h in enumerate(fh)}
    fmt_col = 2  # colonne C : format concerné (sans en-tête dans la grille)
    alu_by_format = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[col["Réf"]] and str(r[col["Réf"]]).startswith("ALU"):
            alu_by_format[str(r[col["Format"]]).replace("x", "×").replace(" ", "")] = r[col["Réf"]]
    for r in wf.iter_rows(min_row=2, values_only=True):
        if not r[fc["Support"]] or not r[fc["Finition"]]:
            continue
        fmt = r[fmt_col]
        pref = alu_by_format.get(str(fmt).replace("x", "×").replace(" ", "")) if fmt else None
        fin.append("(%s,%s,%s,%s,%d,%s)" % (
            q(support_of(r[fc["Support"]])), q(r[fc["Finition"]]), q(r[fc["Type"]]),
            q(pref), cents(r[fc["Supplément HT (€)"]] or 0), q(r[fc["Reference"]])))

    with open(out, "w", encoding="utf-8") as f:
        f.write("-- Généré par scripts/build_seed.py depuis la grille tarifaire. Ne pas éditer à la main.\n\n")
        f.write("insert into public.products (ref, support, variant, format_label, width_cm, height_cm,"
                " price_ttc, cost_ht, min_qty, sort_order) values\n")
        f.write(",\n".join(rows))
        f.write("\non conflict (ref) do update set support=excluded.support, variant=excluded.variant,"
                " format_label=excluded.format_label, width_cm=excluded.width_cm, height_cm=excluded.height_cm,"
                " price_ttc=excluded.price_ttc, cost_ht=excluded.cost_ht, min_qty=excluded.min_qty,"
                " sort_order=excluded.sort_order, active=true;\n\n")
        f.write("delete from public.finishes;\n")
        f.write("insert into public.finishes (support, name, detail, product_ref, price_ht, supplier_ref) values\n")
        f.write(",\n".join(fin) + ";\n")
    print(f"{len(rows)} produits, {len(fin)} finitions -> {out}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
